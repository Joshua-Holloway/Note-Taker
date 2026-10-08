import pyaudio, wave, sys
from pathlib import Path
import pymupdf
from docx import Document
from pptx import Presentation
from faster_whisper import WhisperModel
from PySide6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QPushButton,
    QLabel,
    QPlainTextEdit,
    QComboBox,
    QTabWidget,
    QListWidget,
    QVBoxLayout,
    QHBoxLayout,
    QFrame,
    QFileDialog
)
from PySide6.QtCore import Qt, QObject, QThread, Signal, Slot
from ollama import chat
model_size = "small"
chunk = 1024  # Record in chunks of 1024 samples
sample_format = pyaudio.paInt16  # 16 bits per sample
channels = 1
fs = 44100  # Record at 44100 samples per second
seconds = 10
filename = "output.wav"


class NoteGenerationWorker(QObject):
    finished = Signal(str)
    error = Signal(str)

    def __init__(self, transcription, file):
        self.file = file
        self.transcription = transcription
        super().__init__()

    @Slot()
    def run(self):
        SYSTEM_PROMPT = """
        
You are converting university lecture material into complete revision notes.

Your task is to summarise the ENTIRE supplied lecture faithfully.

Do not turn the lecture into a tutorial, lab guide, coursework solution, or generic study advice.

Do not use your own background knowledge to expand or improve the lecture content.

If a fact, example, historical detail, code pattern, technical rule, or assessment claim is not explicitly supported by the source, omit it.

Prefer a slightly incomplete but faithful note over a more complete note containing outside knowledge.

STRICT RULES:

1. Use only information supported by the supplied source.
2. Do not invent:
   - lab requirements
   - submission instructions
   - deadlines
   - filenames
   - commands
   - assessment advice
   - lecturer intentions
   - examples not present in the source
3. An example is not a requirement.
4. Do not tell the student what they "must do" unless the source explicitly says so.
5. Do not add "next steps", motivational language, offers of help, or generic advice.
6. Cover ALL major academic sections of the lecture, not only the final slides or lab section.
7. Follow the lecture's conceptual order.
8. Give important topics detail proportional to how much attention they receive in the source.
9. Preserve important:
   - definitions
   - examples
   - technical syntax
   - relationships
   - rules
   - comparisons
   - lecturer explanations
10. If information is unclear, incomplete, or missing, say so rather than guessing.
11. Use professional British English.
12. Use clean Markdown and make the notes easy to revise from.

POWERPOINT RULES:
- Treat slide order as meaningful.
- Use slide boundaries to identify topic changes.
- Combine related slides into coherent sections.
- Ignore isolated slide numbers and obvious extraction noise.
- Do not infer unseen diagram content.
- Do not disproportionately focus on slides containing words such as "Lab", "Assessment", "Summary", or "Next Week".

OUTPUT:

# Lecture title

## Lecture overview
Explain the overall subject and how the major topics progress.

## Learning objectives
Only include objectives explicitly stated in the source.

## Detailed lecture notes
Cover the full lecture in order.

Use meaningful headings for each major topic.

For each topic:
- explain the concept clearly,
- preserve important definitions,
- include useful examples from the source,
- explain important technical rules or relationships.

Use code blocks for code, JDL, SQL, formulas, or other technical syntax.

Use tables only when they improve clarity.

## Key definitions
Provide a concise glossary of important terms.

## Key rules and relationships
Include only rules actually taught in the source.

## Lab and assessment information
Include only information explicitly stated in the source.
Do not invent missing requirements.

## Revision checklist
Create short "I can..." statements based only on material actually covered.

## Summary
Give a concise synthesis of the lecture.

FINAL CHECK:
Before answering, make sure:
- every major lecture section is represented,
- no unsupported requirement has been added,
- no example has been turned into a requirement,
- no lab or assessment detail has been invented,
- the output reads like revision notes rather than advice.

If a statement is not supported by the source, remove it.
"""
        user_prompt = f"""
Create complete revision notes from the lecture below.

Important:
- cover the entire lecture,
- do not focus disproportionately on the lab section,
- do not add information that is not present in the source.

SOURCE:

{self.transcription}
"""
        try:
            response = chat(
                model="qwen3:4b",
                messages=[
                    {
                        "role": "system",
                        "content": SYSTEM_PROMPT
                    },
                    {
                        "role": "user",
                        "content": user_prompt
                    }
                ],
                options={
                    "num_ctx": 8192,
                    "temperature": 0
                }
            )

            with open(self.file, "w", encoding="utf-8") as f:
                f.write(response.message.content)
            self.finished.emit(self.file)

        except Exception as e:
            self.error.emit(str(e))


class TranscribingWorker(QObject):
    finished = Signal(str)
    error = Signal(str)

    def __init__(self, input_file, output_file, whisper_model):
        # Keep the media file, output file and selected Whisper model separate.
        self.input_file = input_file
        self.output_file = output_file
        self.whisper_model = whisper_model
        super().__init__()

    @Slot()
    def run(self):
        try:
            print("Started transcribing")
            model = WhisperModel(
                self.whisper_model,
                device="cpu",
                compute_type="int8"
            )

            # faster-whisper can read the selected audio/video file directly.
            segments, info = model.transcribe(
                self.input_file,
                beam_size=5
            )

            # Save the transcript so I still have a text copy outside the GUI.
            with open(self.output_file, "w", encoding="utf-8") as f:
                for segment in segments:
                    f.write(f"{segment.text}\n")
            print("Finished transcribing")
            self.finished.emit(self.output_file)

        except Exception as e:
            self.error.emit(str(e))


class RecordingWorker(QObject):
    finished = Signal(str)
    error = Signal(str)

    def __init__(self):
        super().__init__()

    @Slot()
    def run(self):
        p = None
        stream = None
        frames = []

        try:
            p = pyaudio.PyAudio()
            stream = p.open(
                format=sample_format,
                channels=channels,
                rate=fs,
                frames_per_buffer=chunk,
                input=True
            )
            print("Recording")
            while not QThread.currentThread().isInterruptionRequested():
                data = stream.read(chunk)
                frames.append(data)
            print("Recording finished")
            stream.stop_stream()
            stream.close()
            stream = None
            wf = wave.open(filename, "wb")
            wf.setnchannels(channels)
            wf.setsampwidth(p.get_sample_size(sample_format))
            wf.setframerate(fs)
            wf.writeframes(b"".join(frames))
            wf.close()
            self.finished.emit(filename)

        except Exception as e:
            self.error.emit(str(e))

        finally:
            if stream is not None:
                stream.stop_stream()
                stream.close()

            if p is not None:
                p.terminate()


class MainWindow(QMainWindow):

    def __init__(self):
        super().__init__()
        self.record_thread = None
        self.transcribing_thread = None
        self.generating_thread = None
        self.recording_worker = None
        self.transcribing_worker = None
        self.generating_worker = None

        # Keep track of the selected source file and the text extracted from it.
        self.current_file = None
        self.current_text = None
        self.setWindowTitle("Lecture Assistant")
        self.resize(1200, 800)

        # --------------------------------------------------
        # CENTRAL WIDGET
        # --------------------------------------------------
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QHBoxLayout(central_widget)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # --------------------------------------------------
        # SIDEBAR
        # --------------------------------------------------
        sidebar = QFrame()
        sidebar.setFixedWidth(220)
        sidebar_layout = QVBoxLayout(sidebar)
        sidebar_layout.setContentsMargins(20, 25, 20, 20)
        sidebar_layout.setSpacing(12)
        app_title = QLabel("Lecture Assistant")
        app_title.setObjectName("appTitle")
        sidebar_layout.addWidget(app_title)

        # This will eventually contain the lectures I have saved.
        self.lecture_list = QListWidget()

        # Temporary examples so the sidebar is not empty while I build the lecture library.
        self.lecture_list.addItem("Databases and Web Development")
        self.lecture_list.addItem("Functional Programming")
        self.lecture_list.addItem("Artificial Intelligence 2")
        sidebar_layout.addWidget(self.lecture_list)
        new_lecture_button = QPushButton("+ New Lecture")

        # This will reset the interface for a new lecture later.
        sidebar_layout.addWidget(new_lecture_button)
        main_layout.addWidget(sidebar)

        # --------------------------------------------------
        # MAIN CONTENT AREA
        # --------------------------------------------------
        content = QWidget()
        content_layout = QVBoxLayout(content)
        content_layout.setContentsMargins(30, 25, 30, 25)
        content_layout.setSpacing(18)
        main_layout.addWidget(content)

        # --------------------------------------------------
        # HEADER
        # --------------------------------------------------
        header_layout = QHBoxLayout()
        title_area = QVBoxLayout()
        self.lecture_title = QLabel("New Lecture")
        self.lecture_title.setObjectName("lectureTitle")
        self.lecture_subtitle = QLabel(
            "Start recording or import an existing lecture."
        )
        title_area.addWidget(self.lecture_title)
        title_area.addWidget(self.lecture_subtitle)
        header_layout.addLayout(title_area)

        # Keep the model controls aligned to the right of the header.
        header_layout.addStretch()

        # --------------------------------------------------
        # WHISPER MODEL SELECTOR
        # --------------------------------------------------
        model_label = QLabel("Whisper model:")
        self.model_selector = QComboBox()
        self.model_selector.addItems([
            "tiny",
            "base",
            "small",
            "medium",
            "large-v3"
        ])
        self.model_selector.setCurrentText(model_size)
        header_layout.addWidget(model_label)
        header_layout.addWidget(self.model_selector)
        content_layout.addLayout(header_layout)

        # --------------------------------------------------
        # RECORDING CONTROLS
        # --------------------------------------------------
        controls = QFrame()
        controls_layout = QHBoxLayout(controls)
        self.record_button = QPushButton("● Record")
        self.stop_button = QPushButton("■ Stop")
        self.import_button = QPushButton("Import File")
        self.transcribe_button = QPushButton("Transcribe")
        self.generate_notes_button = QPushButton("Generate Notes")
        self.stop_button.setEnabled(False)

        # --------------------------------------------------
        # BUTTON FUNCTIONS GO HERE
        # --------------------------------------------------
        self.record_button.clicked.connect(self.start_recording)
        self.stop_button.clicked.connect(self.stop_recording)
        self.transcribe_button.clicked.connect(self.start_transcribing)
        self.generate_notes_button.clicked.connect(self.start_generating_notes)
        self.import_button.clicked.connect(self.choose_file)

        # Drag-and-drop can use the same file handling later.
        # Eventually this could open QFileDialog
        # OR become your drag-and-drop alternative.
        # self.transcribe_button.clicked.connect(...)
        # self.generate_notes_button.clicked.connect(...)
        #
        # Eventually this will call Ollama note generation.
        controls_layout.addWidget(self.record_button)
        controls_layout.addWidget(self.stop_button)
        controls_layout.addSpacing(20)
        controls_layout.addWidget(self.import_button)
        controls_layout.addWidget(self.transcribe_button)
        controls_layout.addStretch()
        controls_layout.addWidget(self.generate_notes_button)
        content_layout.addWidget(controls)

        # --------------------------------------------------
        # STATUS BAR AREA
        # --------------------------------------------------
        status_layout = QHBoxLayout()
        self.status_label = QLabel("● Idle")
        self.duration_label = QLabel("00:00:00") #doesnt do anything yet
        status_layout.addWidget(self.status_label)
        status_layout.addStretch()
        status_layout.addWidget(self.duration_label)
        content_layout.addLayout(status_layout)

        # --------------------------------------------------
        # MAIN LECTURE TABS
        # --------------------------------------------------
        self.tabs = QTabWidget()

        # -------------------
        # TRANSCRIPT TAB
        # -------------------
        transcript_page = QWidget()
        transcript_layout = QVBoxLayout(transcript_page)
        self.transcript_box = QPlainTextEdit()
        self.transcript_box.setPlaceholderText(
            "Lecture transcript will appear here..."
        )
        transcript_layout.addWidget(self.transcript_box)
        self.tabs.addTab(transcript_page, "Transcript")

        # -------------------
        # NOTES TAB
        # -------------------
        notes_page = QWidget()
        notes_layout = QVBoxLayout(notes_page)
        self.notes_box = QPlainTextEdit()
        self.notes_box.setPlaceholderText(
            "AI-generated lecture notes will appear here..."
        )
        notes_layout.addWidget(self.notes_box)
        self.tabs.addTab(notes_page, "Notes")

        # -------------------
        # SUMMARY TAB
        # -------------------
        summary_page = QWidget()
        summary_layout = QVBoxLayout(summary_page)
        self.summary_box = QPlainTextEdit()
        self.summary_box.setPlaceholderText(
            "Lecture summary will appear here..."
        )
        summary_layout.addWidget(self.summary_box)
        self.tabs.addTab(summary_page, "Summary")
        content_layout.addWidget(self.tabs)

        # --------------------------------------------------
        # BASIC TEMPORARY STYLING
        #
        # Eventually move all of this into
        # a separate style.qss file.
        # --------------------------------------------------
        self.setStyleSheet("""
            QMainWindow {
                background-color: #111318;
            }
            QWidget {
                color: #e8e8ea;
                font-size: 14px;
            }
            QFrame {
                background-color: #171a21;
            }
            QLabel#appTitle {
                font-size: 20px;
                font-weight: 700;
                padding-bottom: 10px;
            }
            QLabel#lectureTitle {
                font-size: 24px;
                font-weight: 700;
            }
            QPushButton {
                background-color: #242832;
                border: 1px solid #343945;
                border-radius: 7px;
                padding: 9px 14px;
            }
            QPushButton:hover {
                background-color: #2d323d;
            }
            QPushButton:disabled {
                color: #666a73;
                background-color: #1b1e24;
            }
            QPlainTextEdit {
                background-color: #171a21;
                border: 1px solid #2d323d;
                border-radius: 8px;
                padding: 12px;
                font-size: 14px;
            }
            QListWidget {
                background-color: transparent;
                border: none;
            }
            QListWidget::item {
                padding: 9px;
                border-radius: 6px;
            }
            QListWidget::item:selected {
                background-color: #292e39;
            }
            QComboBox {
                background-color: #242832;
                border: 1px solid #343945;
                border-radius: 6px;
                padding: 7px;
            }
            QTabWidget::pane {
                border: none;
            }
            QTabBar::tab {
                padding: 10px 18px;
            }
            QTabBar::tab:selected {
                border-bottom: 2px solid #e8e8ea;
            }
        """)

    def choose_file(self):
        # Let me choose either lecture media or a document from the normal file picker.
        file_path, selected_filter = QFileDialog.getOpenFileName(
            self,
            "Choose a lecture file",
            "",
            "All Supported Files (*.wav *.mp3 *.m4a *.mp4 *.webm *.mkv *.pdf *.docx *.pptx *.txt);;"
            "Audio and Video (*.wav *.mp3 *.m4a *.mp4 *.webm *.mkv);;"
            "Documents (*.pdf *.docx *.pptx *.txt);;"
            "All Files (*)"
        )

        # If I cancel the picker, there is nothing else to do.
        if not file_path:
            return
        self.current_file = file_path

        extension = Path(file_path).suffix.lower()
        media_types = {".wav", ".mp3", ".m4a", ".mp4", ".webm", ".mkv"}
        document_types = {".pdf", ".docx", ".pptx", ".txt"}

        if extension in media_types:
            # Media still needs to go through Whisper before I can generate notes from it.
            self.current_text = None
            self.transcript_box.clear()
            self.status_label.setText("● Media selected - ready to transcribe")
            self.lecture_subtitle.setText(Path(file_path).name)

        elif extension in document_types:
            try:
                # Documents already contain text, so I can extract it without Whisper.
                self.current_text = self.extract_document_text(file_path, extension)
                self.transcript_box.setPlainText(self.current_text)
                self.status_label.setText("● Document loaded - ready to generate notes")
                self.lecture_subtitle.setText(Path(file_path).name)

            except Exception as e:
                self.current_text = None
                self.status_label.setText("● File import error")
                print(f"File import error: {e}")

        else:
            self.current_text = None
            self.status_label.setText("● Unsupported file type")

    def extract_document_text(self, file_path, extension):
        # Convert each supported document type into one plain text string for Ollama.
        if extension == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

        if extension == ".pdf":
            text = []

            with pymupdf.open(file_path) as document:
                for page in document:
                    text.append(page.get_text())
            return "\n".join(text)

        if extension == ".docx":
            document = Document(file_path)
            return "\n".join(
                paragraph.text for paragraph in document.paragraphs
                if paragraph.text.strip()
            )

        if extension == ".pptx":
            presentation = Presentation(file_path)
            slides_text = []

            # Keep the slide order so the notes follow the presentation properly.
            for slide_number, slide in enumerate(presentation.slides, start=1):
                slide_text = [f"Slide {slide_number}"]

                for shape in slide.shapes:
                    # Most slide content, including titles and text boxes, is stored as text.
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())

                    # Tables need to be read cell by cell instead.
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_text:
                                slide_text.append(" | ".join(row_text))

                slides_text.append("\n".join(slide_text))

            return "\n\n".join(slides_text)

        raise ValueError(f"Unsupported document type: {extension}")

    def start_generating_notes(self):
        # Both documents and transcripts end up in current_text before reaching Ollama.
        if not self.current_text:
            self.status_label.setText("● No text available for note generation")
            return
        self.generate_notes_button.setEnabled(False)
        self.status_label.setText("҉ Generating")

        self.generating_thread = QThread()
        self.generating_worker = NoteGenerationWorker(self.current_text, "note_test.txt")
        self.generating_worker.moveToThread(self.generating_thread)

        self.generating_thread.started.connect(
            self.generating_worker.run
        )
        self.generating_worker.finished.connect(
            self.generating_finished
        )
        self.generating_worker.error.connect(
            self.generating_error
        )
        self.generating_worker.finished.connect(
            self.generating_thread.quit
        )
        self.generating_worker.finished.connect(
            self.generating_worker.deleteLater
        )
        self.generating_thread.finished.connect(
            self.generating_thread.deleteLater
        )
        self.generating_thread.finished.connect(
            self.generating_thread_finished
        )
        self.generating_worker.error.connect(
            self.generating_thread.quit
        )
        self.generating_worker.error.connect(
            self.generating_worker.deleteLater
        )

        self.generating_thread.start()

    @Slot()
    def generating_thread_finished(self):
        self.generating_thread = None
        self.generating_worker = None

    @Slot(str)
    def generating_finished(self, generated_file):
        # Show the saved Ollama output in the Notes tab once generation is complete.
        with open(generated_file, "r", encoding="utf-8") as f:
            generated_notes = f.read()
        self.notes_box.setPlainText(generated_notes)
        self.tabs.setCurrentWidget(self.notes_box.parentWidget())
        self.status_label.setText("● Notes generated")
        self.generate_notes_button.setEnabled(True)
        print(f"Saved generated notes to {generated_file}")

    @Slot(str)
    def generating_error(self, error_message):
        self.status_label.setText("● Generating error")
        self.generate_notes_button.setEnabled(True)
        print(f"Generating error: {error_message}")

    def start_transcribing(self):
        # Use an imported media file if I have one; otherwise use the latest recording.
        source_file = self.current_file if self.current_file else filename
        extension = Path(source_file).suffix.lower()
        media_types = {".wav", ".mp3", ".m4a", ".mp4", ".webm", ".mkv"}

        if extension not in media_types:
            self.status_label.setText("● Selected file does not need transcription")
            return

        self.transcribe_button.setEnabled(False)
        self.status_label.setText("҉ Transcribing")

        # Use the model currently selected in the GUI rather than always using the default.
        selected_model = self.model_selector.currentText()

        self.transcribing_thread = QThread()
        self.transcribing_worker = TranscribingWorker(source_file, "test.txt", selected_model)
        self.transcribing_worker.moveToThread(self.transcribing_thread)

        self.transcribing_thread.started.connect(
            self.transcribing_worker.run
        )
        self.transcribing_worker.finished.connect(
            self.transcribing_finished
        )
        self.transcribing_worker.error.connect(
            self.transcribing_error
        )
        self.transcribing_worker.finished.connect(
            self.transcribing_thread.quit
        )
        self.transcribing_worker.finished.connect(
            self.transcribing_worker.deleteLater
        )
        self.transcribing_thread.finished.connect(
            self.transcribing_thread.deleteLater
        )
        self.transcribing_thread.finished.connect(
            self.transcribing_thread_finished
        )
        self.transcribing_worker.error.connect(
            self.transcribing_thread.quit
        )
        self.transcribing_worker.error.connect(
            self.transcribing_worker.deleteLater
        )

        self.transcribing_thread.start()

    @Slot()
    def transcribing_thread_finished(self):
        self.transcribing_thread = None
        self.transcribing_worker = None

    @Slot(str)
    def transcribing_finished(self, transcribed_file):
        # Load the completed transcript into the GUI and keep it ready for note generation.
        with open(transcribed_file, "r", encoding="utf-8") as f:
            self.current_text = f.read()
        self.transcript_box.setPlainText(self.current_text)
        self.status_label.setText("● Transcription saved - ready to generate notes")
        self.transcribe_button.setEnabled(True)
        print(f"Saved transcription to {transcribed_file}")

    @Slot(str)
    def transcribing_error(self, error_message):
        self.status_label.setText("● Transcribing error")
        self.transcribe_button.setEnabled(True)
        print(f"Transcribing error: {error_message}")

    @Slot(str)
    def recording_finished(self, recorded_file):
        # Make the latest recording the current media file so it can be transcribed next.
        self.current_file = recorded_file
        self.current_text = None
        self.status_label.setText("● Recording saved")
        self.record_button.setEnabled(True)
        self.stop_button.setEnabled(False)
        print(f"Saved recording to {recorded_file}")

    @Slot(str)
    def recording_error(self, error_message):
        self.status_label.setText("● Recording error")
        self.record_button.setEnabled(True)
        self.stop_button.setEnabled(False)
        print(f"Recording error: {error_message}")

    @Slot()
    def recording_thread_finished(self):
        self.record_thread = None
        self.recording_worker = None

    def start_recording(self):
        self.record_button.setEnabled(False)
        self.stop_button.setEnabled(True)
        self.status_label.setText("● Recording")

        self.record_thread = QThread()
        self.recording_worker = RecordingWorker()
        self.recording_worker.moveToThread(self.record_thread)

        self.record_thread.started.connect(
            self.recording_worker.run
        )
        self.recording_worker.finished.connect(
            self.recording_finished
        )
        self.recording_worker.error.connect(
            self.recording_error
        )
        self.recording_worker.finished.connect(
            self.record_thread.quit
        )
        self.recording_worker.finished.connect(
            self.recording_worker.deleteLater
        )
        self.record_thread.finished.connect(
            self.record_thread.deleteLater
        )
        self.record_thread.finished.connect(
            self.recording_thread_finished
        )
        self.recording_worker.error.connect(
            self.record_thread.quit
        )
        self.recording_worker.error.connect(
            self.recording_worker.deleteLater
        )

        self.record_thread.start()

    def stop_recording(self):
        if self.record_thread is not None:
            self.status_label.setText("● Stopping...")
            self.stop_button.setEnabled(False)
            self.record_thread.requestInterruption()




# ==========================================================
# APPLICATION START
# ==========================================================


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
